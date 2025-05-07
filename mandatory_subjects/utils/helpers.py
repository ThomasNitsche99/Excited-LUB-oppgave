import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


#Function for converting csv to excel
def csv_to_excel(csv_file):
    # Step 1: Read the CSV
    df = pd.read_csv(csv_file)
    
    # Step 2: Process the dataframe to add empty rows between schools
    processed_rows = []
    current_school = None
    
    for idx, row in df.iterrows():
        # Add multiple empty rows when school changes
        if current_school is not None and current_school != row['Skole'] and '---' not in row['Skole']:
            # Add 5 rows between schools: 2 empty, 1 red, 2 empty
            for i in range(5):
                if i == 2:  # Third row (middle)
                    processed_rows.append(['RED_SEPARATOR', '', '', ''])  # Special marker for red row
                else:
                    processed_rows.append(['' for _ in range(len(df.columns))])
        
        processed_rows.append(row.tolist())
        if '---' not in row['Skole']:
            current_school = row['Skole']
    
    # Create new dataframe with added rows
    df_processed = pd.DataFrame(processed_rows, columns=df.columns)

    # Step 3: Write to Excel
    excel_path = f"{csv_file.split('.')[0]}.xlsx"
    df_processed.to_excel(excel_path, index=False)

    # Step 4: Open the workbook and access the worksheet
    wb = load_workbook(excel_path)
    ws = wb.active

    # Step 5: Define styles
    header_font = Font()  # Regular font, no bold
    header_fill = PatternFill("solid", fgColor="F2F2F2")  # Light gray background
    alignment_center = Alignment(horizontal="center", vertical="center")
    alignment_wrap = Alignment(wrap_text=True, vertical="top")

    # Color fills
    fill_kunnskaper = PatternFill("solid", fgColor="DCE6F1")      # light blue
    fill_ferdigheter = PatternFill("solid", fgColor="FFF2CC")     # light yellow/orange
    fill_generell = PatternFill("solid", fgColor="C6EFCE")        # light green
    fill_separator = PatternFill("solid", fgColor="000000")       # black for separator rows
    fill_red_separator = PatternFill("solid", fgColor="FF0000")   # red for separator between schools

    # Step 6: Apply formatting to header and columns
    for col_num, cell in enumerate(ws[1], start=1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment_center

        # Column letter
        col_letter = get_column_letter(col_num)

        # Set column width
        max_length = max((len(str(c.value)) for c in ws[col_letter]), default=10)
        ws.column_dimensions[col_letter].width = max_length + 2

    # Step 7: Apply formatting to data rows
    for row_num, row in enumerate(ws.iter_rows(min_row=2), start=2):
        first_cell_value = str(row[0].value)
        
        # Check if this is a separator row (contains '---')
        if '---' in first_cell_value:
            for cell in row:
                cell.fill = fill_separator
            continue

        # Check if this is a red separator row
        if first_cell_value == 'RED_SEPARATOR':
            for cell in row:
                cell.fill = fill_red_separator
                cell.value = ''  # Clear the marker text
            continue

        # Skip empty rows (they separate schools)
        if not first_cell_value:
            continue

        # Get the learning outcome type cell (third column)
        type_cell = row[2]
        learning_type = str(type_cell.value).lower() if type_cell.value else ""

        # Apply color based on learning outcome type to only the type column
        if "kunnskaper" in learning_type:
            type_cell.fill = fill_kunnskaper
        elif "ferdigheter" in learning_type:
            type_cell.fill = fill_ferdigheter
        elif "generell" in learning_type:
            type_cell.fill = fill_generell
                
        # Apply text wrapping to the last column
        row[-1].alignment = alignment_wrap

    # Optional: Freeze top row
    ws.freeze_panes = "A2"

    # Save changes
    wb.save(excel_path)
    print(f"\nExcel file saved: {excel_path}")


# Function for creating csv file
def create_csv(school: str, study_program: str, LUBs: dict):
    rows = []
    for category, descriptions in LUBs.items():
        for description in descriptions:
            rows.append([school, study_program, category.capitalize(), description])

    df = pd.DataFrame(rows, columns=["Skole", "Studie program", "Læringsutbytte type", "Læringsutbytte"])
    separator_row = pd.DataFrame([["-"*6, "-"*30, "-"*10, "-"*200]], columns=df.columns)
    df = pd.concat([df, separator_row], ignore_index=True)
    return df

